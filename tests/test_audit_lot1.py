"""
Tests de non-régression du lot 1 de l'audit :
lot de fichiers robuste, scaler DLMS, nombres localisés, NaN, n° de compteur, ZIP homonymes.

Lancer depuis la racine du projet : python -m pytest tests
"""
import io
import sys
import zipfile
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import Manual_Meter_Readings as m  # noqa: E402
from test_timezone import build_lp2_csv, build_profile_xml, dlms_ts, export_timestamps  # noqa: E402


def export_values(data: m.ParsedMeterData):
    document, _ = m.convert_to_json(data, "Landis", "LGZ")
    return [
        block["IntervalReadings"][0]["value"]
        for reading in document["payload"]["MeterReadings"]
        for block in reading["IntervalBlocks"]
    ]


class Upload(io.BytesIO):
    """Imite un UploadedFile Streamlit."""

    def __init__(self, name: str, content: bytes):
        super().__init__(content)
        self.name = name


@pytest.fixture
def session(monkeypatch):
    state = SimpleNamespace(user_configs={})
    monkeypatch.setattr(m.st, "session_state", state)
    return state


def build_zip(files) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, content in files:
            archive.writestr(name, content)
    return buffer.getvalue()


CSV_OK = build_lp2_csv([
    ("25/11/2025 00:00:00", "0", "1"),
    ("26/11/2025 00:00:00", "0", "2"),
])


# ============================================================================
# Nombres (G, I)
# ============================================================================

@pytest.mark.parametrize("raw, expected", [
    ("1234.5", 1234.5),
    ("1234,5", 1234.5),
    ("1'234,5", 1234.5),
    ("1’234.5", 1234.5),
    ("1 234,5", 1234.5),
    ("1.234,5", 1234.5),
    ("1,234.5", 1234.5),
    ("1.234.567", 1234567.0),
    ("-12,5", -12.5),
    (42, 42.0),
    (1.5, 1.5),
    ("", None),
    ("abc", None),
    ("NaN", None),
    ("inf", None),
    (float("nan"), None),
    (None, None),
    (True, None),
])
def test_parse_number(raw, expected):
    assert m.parse_number(raw) == expected


def test_csv_localized_numbers_and_invalid_values():
    csv = build_lp2_csv([
        ("25/11/2025 00:00:00", "0", "1'234,5"),
        ("26/11/2025 00:00:00", "0", "abc"),
        ("27/11/2025 00:00:00", "0", "NaN"),
        ("28/11/2025 00:00:00", "0", "1.000,25"),
    ])
    data = m.parse_csv(csv, "12365275-LP2.csv")[0]

    assert export_values(data) == ["1234500", "1000250"]  # kWh → Wh
    assert any("2 valeur(s) non numérique(s)" in w for w in data.warnings)


def test_convert_skips_non_finite_values():
    data = m.ParsedMeterData(
        meter_id="12365275",
        load_profile="Profil de charge 2",
        interval="24h",
        channels={"1-0:1.8.1": {"unit": "Wh", "readings": [
            (datetime(2026, 1, 1), float("nan"), 0),
            (datetime(2026, 1, 2), 5.0, 0),
        ]}},
        source_file="x.csv",
    )
    assert export_values(data) == ["5"]


# ============================================================================
# Scaler DLMS (B, C, T)
# ============================================================================

@pytest.mark.parametrize("quantity, scaler, expected", [
    ("ActiveEnergy", 0, "Wh"),
    ("ActiveEnergy", 3, "kWh"),
    ("ActiveEnergy", -3, "mWh"),
    ("ReactiveEnergy", 6, "Mvarh"),
    ("ActiveEnergy", 2, "Wh*10^2"),
    ("ActiveEnergy", -1, "Wh*10^-1"),
    ("Unknown", 0, ""),
])
def test_resolve_unit_from_scaler(quantity, scaler, expected):
    assert m.resolve_unit_from_scaler(quantity, scaler) == expected


@pytest.mark.parametrize("unit, expected", [
    ("Wh", 1), ("varh", 1), ("", 1), ("inconnu", 1),
    ("kWh", 1000), ("KWH", 1000), ("kvarh", 1000), ("kVAh", 1000), ("kW", 1000),
    ("MWh", 1e6), ("mWh", 1e-3),
    ("Wh*10^2", 100), ("Wh*10^-1", 0.1), ("kWh*10^-1", 100),
])
def test_unit_to_base_multiplier(unit, expected):
    assert m.unit_to_base_multiplier(unit) == pytest.approx(expected)


def build_xml_with_register_scaler(scaler: int) -> bytes:
    xml = build_profile_xml(2, [
        (dlms_ts(datetime(2025, 12, 10)), "0", "1234"),
    ]).decode("utf-8")
    register = (
        '<Objects ObjectName="DD.EnergyRegister_A_import_T1" ObjectLogicalName="0100010801FF" ClassID="3">'
        '<Attributes AttributeName="scaler_unit">'
        f'<Field FieldName="DD.EnergyRegister_A_import_T1.scaler_unit.0.Scaler" FieldValue="{scaler}" />'
        '<Field FieldName="DD.EnergyRegister_A_import_T1.scaler_unit.0.Quantity" FieldValue="ActiveEnergy" />'
        "</Attributes></Objects>"
    )
    return xml.replace("</Root>", register + "</Root>").encode("utf-8")


@pytest.mark.parametrize("scaler, unit, value", [
    (0, "Wh", "1234"),
    (3, "kWh", "1234000"),
    (-1, "Wh*10^-1", "123"),
])
def test_xml_register_scaler_applied(scaler, unit, value):
    data = m.parse_xml(build_xml_with_register_scaler(scaler), "scaler.xml")[0]

    assert data.channels["1-0:1.8.1"]["unit"] == unit
    assert export_values(data) == [value]


def test_obis_handler_parses_obis_groups():
    assert m.OBISCodeHandler.infer_unit_from_obis("1-0:1.8.0") == "Wh"
    assert m.OBISCodeHandler.infer_unit_from_obis("1-0:5.8.1") == "varh"
    assert m.OBISCodeHandler.infer_unit_from_obis("1-0:1.8.0", scaler=3) == "kWh"
    assert m.OBISCodeHandler.is_energy_obis("1-0:1.8.0")
    assert not m.OBISCodeHandler.is_energy_obis("0-0:96.10.1")
    assert m.OBISCodeHandler.get_display_name("1-0:1.8.0") == "1.8.0"


# ============================================================================
# Traitement du lot (A, J, K)
# ============================================================================

def test_unrecognized_xml_does_not_abort_batch(session):
    archive = build_zip([("bad.xml", "<Root/>"), ("12365275-LP2.csv", CSV_OK)])
    data, warnings, needing_config = m.process_uploaded_files(
        [Upload("lot.zip", archive)], ";", "Landis", "ESR10307"
    )

    assert [d.meter_id for d in data] == ["12365275"]
    assert any(w.startswith("lot.zip/bad.xml: fichier ignoré") for w in warnings)
    assert not needing_config


def test_unreadable_file_reported(session):
    data, warnings, _ = m.process_uploaded_files(
        [Upload("ancien.xls", b"\xd0\xcf\x11\xe0 binaire"), Upload("12365275-LP2.csv", CSV_OK)],
        ";", "Landis", "ESR10307"
    )

    assert len(data) == 1
    assert any(w.startswith("ancien.xls: aucune donnée exploitable") for w in warnings)


def build_xlsx_without_serial() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "Informations"
    sheet = wb.create_sheet("Données")
    sheet.append(["Horloge 0-0:1.0.0", "1-0:1.8.0 (kWh)"])
    sheet.append([datetime(2026, 1, 1, 0, 0), 1])
    sheet.append([datetime(2026, 1, 1, 0, 15), 2])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_xlsx_serial_from_filename():
    data = m.parse_xlsx(build_xlsx_without_serial(), "12345678-20260101.xlsx")[0]

    assert data.meter_id == "12345678"
    assert any("déduit du nom de fichier" in w for w in data.warnings)


def test_xlsx_without_serial_not_exported(session):
    data, warnings, _ = m.process_uploaded_files(
        [Upload("releve.xlsx", build_xlsx_without_serial())], ";", "Landis", "ESR10307"
    )

    assert data == []
    assert any("numéro de compteur introuvable, fichier non exporté" in w for w in warnings)


def test_zip_homonyms_kept_apart(session):
    other_csv = CSV_OK.replace(b"12365275", b"87654321")
    archive = build_zip([("a/releve.csv", CSV_OK), ("b/releve.csv", other_csv)])
    data, _, _ = m.process_uploaded_files([Upload("lot.zip", archive)], ";", "Landis", "ESR10307")

    assert sorted(d.meter_id for d in data) == ["12365275", "87654321"]
    assert sorted(d.source_file for d in data) == ["lot.zip/a/releve.csv", "lot.zip/b/releve.csv"]


def test_manual_config_key_depends_on_content():
    assert m.file_config_key("releve.csv", b"a") != m.file_config_key("releve.csv", b"b")
    assert m.file_config_key("releve.csv", b"a") == m.file_config_key("releve.csv", b"a")


UNKNOWN_CSV = b"releve sans entete\nHorloge;Valeur;Autre\n01/01/2026 00:00;1.5;2.5\n01/01/2026 00:15;3.5;4.5\n"
MANUAL_CONFIG = {"meter_id": "11223344", "header_line": 1, "data_start": 2, "interval": "15min"}


def test_manual_config_applied_by_content_key(session):
    _, _, needing_config = m.process_uploaded_files(
        [Upload("inconnu.csv", UNKNOWN_CSV)], ";", "Landis", "ESR10307"
    )
    key = m.file_config_key("inconnu.csv", UNKNOWN_CSV)
    assert list(needing_config) == [key]
    assert needing_config[key]["name"] == "inconnu.csv"

    session.user_configs[key] = MANUAL_CONFIG
    data, _, needing_config = m.process_uploaded_files(
        [Upload("inconnu.csv", UNKNOWN_CSV)], ";", "Landis", "ESR10307"
    )
    assert not needing_config
    assert [d.meter_id for d in data] == ["11223344"]

    # Même nom, autre contenu : la configuration mémorisée ne s'applique pas
    _, _, needing_config = m.process_uploaded_files(
        [Upload("inconnu.csv", UNKNOWN_CSV.replace(b"1.5;2.5", b"5.5;6.5"))], ";", "Landis", "ESR10307"
    )
    assert len(needing_config) == 1


def test_simultaneous_homonyms_needing_config_kept_apart(session):
    other = UNKNOWN_CSV.replace(b"1.5;2.5", b"5.5;6.5")
    _, _, needing_config = m.process_uploaded_files(
        [Upload("inconnu.csv", UNKNOWN_CSV), Upload("inconnu.csv", other)], ";", "Landis", "ESR10307"
    )

    assert len(needing_config) == 2
    assert {info["name"] for info in needing_config.values()} == {"inconnu.csv"}


def test_parse_and_collect_manual_config_path():
    """Chemin utilisé par la validation manuelle dans main()."""
    data, warnings, still_needing = [], [], {}
    m.parse_and_collect("inconnu.csv", UNKNOWN_CSV, ";", MANUAL_CONFIG, data, warnings, still_needing)

    assert not still_needing
    assert [d.meter_id for d in data] == ["11223344"]
    outputs, _ = m.build_json_outputs(data, "Landis", "ESR10307")
    assert outputs == []  # colonnes sans OBIS : rien d'exportable, pas de faux succès


def test_file_without_readings_not_exported(session):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "Informations"
    wb.active.append(["Numéro de série", "12365275"])
    sheet = wb.create_sheet("Données")
    sheet.append(["Horloge 0-0:1.0.0", "1-0:1.8.0 (kWh)"])
    sheet.append([datetime(2026, 1, 1, 0, 0), "abc"])
    sheet.append([datetime(2026, 1, 1, 0, 15), "NaN"])
    buffer = io.BytesIO()
    wb.save(buffer)

    data, warnings, _ = m.process_uploaded_files([Upload("vide.xlsx", buffer.getvalue())], ";", "Landis", "ESR10307")

    assert data == []
    assert any("2 valeur(s) non numérique(s)" in w for w in warnings)
    assert any("vide.xlsx: aucun relevé exploitable, fichier non exporté" in w for w in warnings)


def test_encrypted_zip_member_does_not_abort_batch(session):
    archive = bytearray(build_zip([("chiffre.csv", b"x"), ("12365275-LP2.csv", CSV_OK)]))
    # Marque le premier membre comme chiffré (bit 0 des flags, en-tête local et répertoire central)
    local = archive.find(b"PK\x03\x04")
    central = archive.find(b"PK\x01\x02")
    archive[local + 6] |= 1
    archive[central + 8] |= 1

    data, warnings, _ = m.process_uploaded_files([Upload("lot.zip", bytes(archive))], ";", "Landis", "ESR10307")

    assert [d.meter_id for d in data] == ["12365275"]
    assert any("lot.zip/chiffre.csv: illisible dans le ZIP" in w for w in warnings)


def test_export_skips_overflowing_values():
    data = m.ParsedMeterData(
        meter_id="12365275",
        load_profile="Profil de charge 2",
        interval="24h",
        channels={"1-0:1.8.1": {"unit": "kWh", "readings": [
            (datetime(2026, 1, 1), 1e308, 0),
            (datetime(2026, 1, 2), 2.5, 0),
        ]}},
        source_file="x.csv",
    )
    assert export_values(data) == ["2500"]


@pytest.mark.parametrize("raw, expected", [
    ("1,23,4", None),
    ("12.34.5", None),
    ("1.23,5", None),
    ("1,234,567", 1234567.0),
    ("-1.234,5", -1234.5),
    ("-1,25e-3", -0.00125),
    ("1,234", 1.234),  # virgule seule = décimale (convention des fichiers suisses)
])
def test_parse_number_grouping(raw, expected):
    assert m.parse_number(raw) == expected
